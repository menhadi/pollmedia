"""Extract selected A-02 years without losing formulas, cells or footnotes."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import sys
from types import SimpleNamespace

import openpyxl


def workbook_rows(path):
    if Path(path).read_bytes()[:8] == bytes.fromhex('d0cf11e0a1b11ae1'):
        sys.path.insert(0, str(Path(__file__).parent/'tmp'/'python-libs'))
        import xlrd
        book = xlrd.open_workbook(path)
        try:
            names = [name for name in book.sheet_names() if name.strip() == 'A-2']
            if len(names) != 1:
                raise ValueError('Expected one A-2 sheet')
            sheet = book.sheet_by_name(names[0])
            for index in range(sheet.nrows):
                cells = [SimpleNamespace(value=c.value, data_type=c.ctype) for c in sheet.row(index)]
                yield cells, cells
        finally:
            book.release_resources()
        return
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = [name for name in book.sheetnames if name.strip() == 'A-2']
        if len(names) != 1:
            raise ValueError('Expected one A-2 sheet')
        yield from zip(book[names[0]].iter_rows(), cached[names[0]].iter_rows())
    finally:
        book.close()
        cached.close()


def extract(path, manifest, years=(1901, 1911)):
    if hashlib.sha256(Path(path).read_bytes()).hexdigest() != manifest['sha256']:
        raise ValueError('Original checksum mismatch')
    if not years or set(years) - {1901, 1911, 1921, 1931, 1941}:
        raise ValueError('This extractor supports explicitly selected 1901, 1911, 1921, 1931 and 1941 years')
    rows = workbook_rows(path)
    try:
        records, notes, raw_rows, seen = [], [], [], set()
        identity = None
        header = None
        for number, (cells, values) in enumerate(rows, 1):
            raw = [c.value for c in cells]
            headers = [re.sub(r'\s+', '', str(v)).lower() for v in raw[:5]]
            if number == 2:
                header = headers
            if number == 3:
                if all(h in ('none', '') for h in header):
                    header = headers
                headers = [(header[i] + headers[i]) if (header[i], headers[i]) in (
                    ('state', 'code'), ('district', 'code'), ('census', 'year'),
                    ('india/state/', 'unionterritory')) else header[i] for i in range(5)]
            if number == 3 and headers not in [
                    ['statecode', 'districtcode', 'state/district', 'censusyear', 'persons'],
                    ['statecode', 'districtcode', 'unionterritory/district', 'censusyear', 'persons'],
                    ['statecode', 'districtcode', 'india/state/unionterritory', 'censusyear', 'persons'],
                    ['state', 'district', 'state/unionterritory/district', 'censusyear', 'persons'],
                    ['statecode', 'districtcode', 'state/unionterritory/district', 'censusyear', 'persons']]:
                raise ValueError('A-02 headers changed')
            if not any(v is not None and str(v).strip() for v in raw):
                continue
            raw_rows.append({'source_row': number, 'cells': raw,
                             'cell_types': [c.data_type for c in cells],
                             'cached_cells': [c.value for c in values]})
            if number < 5:
                continue
            label = str(raw[3] or '').strip()
            match = re.fullmatch(r'([^\w\s]{0,4})\s*(\d{4})(?:\.0)?\s*([^\w\s]{0,4})', label)
            if not match:
                notes.append({'source_row': number, 'cells': raw})
                continue
            if raw[1] is not None and str(raw[1]).strip():
                def code(v):
                    return str(int(v)) if isinstance(v, (int, float)) and v == int(v) else str(v).strip()
                state = code(raw[0]) if raw[0] is not None and str(raw[0]).strip() else (identity[0] if identity else None)
                identity = (state, code(raw[1]), str(raw[2]).strip())
                if not re.fullmatch(r'\d{2}', identity[0]) or not re.fullmatch(r'\d{3}', identity[1]) or not identity[2]:
                    raise ValueError('Invalid printed identity')
            if identity is None:
                raise ValueError('Year precedes printed geography')
            year = int(match[2])
            key = (*identity[:2], year)
            if key in seen:
                raise ValueError('Duplicate geography/year')
            seen.add(key)
            if year not in years:
                continue
            counts, flags = {}, []
            for field, index in [('persons', 4), ('males', 7), ('females', 8)]:
                value = raw[index]
                if isinstance(value, str) and value.startswith('='):
                    value = values[index].value
                    flags.append(field + ' uses a formula cached by the original workbook; formula preserved')
                if isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 0 and value == int(value):
                    counts[field] = int(value)
                else:
                    counts[field] = None
                    flags.append(field + ' is not a numeric count in the original')
            if all(v is not None for v in counts.values()) and counts['persons'] != counts['males'] + counts['females']:
                flags.append('Persons differs from male plus female components')
            marker = match[1] + match[3]
            if marker:
                flags.append('Printed year marker ' + marker + ': see source footnotes')
            records.append(dict(state_code=identity[0], district_code=identity[1], name=identity[2],
                                year=year, year_label=label, source_row=number,
                                flags=flags, **counts))
        if set(years) != {r['year'] for r in records}:
            raise ValueError('Requested year missing')
        return dict(source=manifest, sheet='A-2', records=records, notes=notes, raw_rows=raw_rows,
                    status='pending_review', boundary_basis=manifest['boundary_basis'],
                    formula_evidence=('BIFF formula expressions remain in preserved original; xlrd supplies cached values'
                                      if Path(path).suffix == '.xls' else 'Formula expressions and cached cells preserved'))
    finally:
        rows.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('original', type=Path)
    parser.add_argument('manifest', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    data = extract(args.original, json.loads(args.manifest.read_text()))
    if args.output.exists():
        raise FileExistsError('Preserve existing extraction; choose a new output')
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'records': len(data['records']), 'flagged': sum(bool(r['flags']) for r in data['records'])}))
