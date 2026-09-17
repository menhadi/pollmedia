"""Bounded extraction of public tabular source files; never executes workbook formulas."""
import csv
import datetime
import io
import json
import re
import sys
import zipfile
from pathlib import Path

MAX_ROWS = 20000
MAX_COLUMNS = 100

def cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        raise ValueError("Nested JSON fields need a source-specific mapping.")
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    result = str(value).strip()
    if len(result) > 4000:
        raise ValueError("A cell exceeds the supported length; a source-specific import is required.")
    return result

def table(rows, header_row=1, max_rows=MAX_ROWS, row_filter=None):
    rows = iter(rows)
    for _ in range(header_row - 1):
        next(rows, None)
    headers = [cell(v).replace("\n", " ") for v in next(rows, [])]
    while headers and headers[-1] == "":
        headers.pop()
    if not headers or len(headers) > MAX_COLUMNS or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("Column headers are blank, duplicated or too wide. Check the header row setting.")
    records = []
    if row_filter and row_filter[0] not in headers:
        raise ValueError('Configured filter column was not found.')
    for number, values in enumerate(rows, header_row + 1):
        if number > 100000:
            raise ValueError('Source exceeds the scanned row limit.')
        values = [cell(v) for v in values]
        if not any(values) or values[:len(headers)] == headers:
            continue
        if len(values) > len(headers) and any(values[len(headers):]):
            raise ValueError("Rows do not match the header columns. Review this source layout.")
        values = (values + [""] * len(headers))[:len(headers)]
        if row_filter and values[headers.index(row_filter[0])] != row_filter[1]:
            continue
        records.append(dict(zip(headers, values)))
        if len(records) > max_rows:
            raise ValueError("Source exceeds the configured row limit; configure a bounded source-specific import.")
    return headers, records

def extract(path, kind, options):
    header = int(options.get("header_row", 1))
    if kind == "csv":
        text = path.read_text(encoding="utf-8-sig")
        delimiter = options.get("delimiter", ",")
        headers, rows = table(csv.reader(io.StringIO(text), delimiter=delimiter), header)
        scope = "CSV file"
    elif kind == "json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        root = data
        for key in filter(None, options.get("json_path", "").split(".")):
            data = data[int(key)] if isinstance(data, list) else data[key]
        if not isinstance(data, list) or not data or any(not isinstance(row, dict) for row in data):
            raise ValueError("JSON path must select a non-empty array of records.")
        if isinstance(root, dict):
            total = root.get("total", root.get("total_records", root.get("count")))
            if isinstance(total, int) and total > len(data):
                raise ValueError("API response is paginated. Configure a complete export or a pagination adapter.")
        headers = list(data[0])
        if any(set(row) != set(headers) for row in data):
            raise ValueError("JSON records have inconsistent fields.")
        headers, rows = table([headers] + [[row[h] for h in headers] for row in data])
        scope = "JSON array: " + (options.get("json_path") or "root") + "; one response (pagination is not followed)"
    elif kind == "xls":
        try:
            import xlrd
        except ImportError:
            sys.path.insert(0, str(Path(__file__).resolve().parent / 'tmp' / 'python-libs'))
            try:
                import xlrd
            except ImportError:
                raise ValueError('Legacy XLS extraction requires the existing xlrd runtime dependency.')
        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            sheet = workbook.sheet_by_name(options.get('sheet') or workbook.sheet_names()[0])
            if sheet.ncols > MAX_COLUMNS or sheet.nrows > MAX_ROWS + header:
                raise ValueError('Legacy workbook exceeds the extraction limits.')
            headers, rows = table((sheet.row_values(index) for index in range(sheet.nrows)), header)
        finally:
            workbook.release_resources()
        scope = 'Legacy XLS worksheet: ' + sheet.name + '; stored cell values only; formulas are not evaluated'
    elif kind == "xlsx":
        import openpyxl
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100000000:
                raise ValueError("Expanded workbook exceeds the import limit.")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        sheet_name = options.get("sheet") or workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise ValueError("Worksheet exceeds 100 columns.")
        def values():
            for row in sheet.iter_rows():
                result = []
                for item in row:
                    if item.data_type == "f":
                        raise ValueError("Worksheet contains formulas. Supply an official values-only export for review.")
                    value = item.value
                    if isinstance(value, int) and re.fullmatch(r"0{2,}", item.number_format or ""):
                        value = str(value).zfill(len(item.number_format))
                    result.append(value)
                yield result
        try:
            maximum = int(options.get('max_rows', MAX_ROWS))
            if not 1 <= maximum <= MAX_ROWS:
                raise ValueError('Workbook row limit must be between 1 and 20,000; partition larger sources.')
            selected = (options['filter_column'], options.get('filter_value', '')) if options.get('filter_column') else None
            headers, rows = table(values(), header, maximum, selected)
        finally:
            workbook.close()
        scope = "Worksheet: " + sheet_name + "; header row " + str(header)
    elif kind == "pdf":
        import pdfplumber
        rows = []
        headers = None
        index = int(options.get("table_index", 1)) - 1
        with pdfplumber.open(path) as document:
            if len(document.pages) > 100:
                raise ValueError("PDF exceeds 100 pages; a source-specific extraction is required.")
            for page in document.pages:
                tables = page.extract_tables()
                if index >= len(tables):
                    raise ValueError("No matching table on PDF page " + str(page.page_number) + ". Scanned PDFs or changed layouts require OCR/manual mapping.")
                found_headers, found_rows = table(tables[index], header)
                if headers is not None and headers != found_headers:
                    raise ValueError("PDF table headers change between pages. A source-specific mapping is required.")
                headers = found_headers
                rows.extend(found_rows)
                if len(rows) > MAX_ROWS:
                    raise ValueError("PDF exceeds the row limit.")
            scope = "PDF table " + str(index + 1) + " on each of " + str(len(document.pages)) + " pages; extracted cells require visual review"
    else:
        raise ValueError("Unsupported source format.")
    if options.get("key_columns"):
        columns = options['key_columns']
        if not isinstance(columns, list) or not columns or any(column not in headers for column in columns) or 'source_record_key' in headers:
            raise ValueError('Invalid composite source key columns.')
        headers.append('source_record_key')
        for row in rows:
            row['source_record_key'] = json.dumps([row[column] for column in columns], ensure_ascii=False, separators=(',', ':'))
    if options.get("filter_column"):
        column = options["filter_column"]
        if column not in headers:
            raise ValueError("Configured filter column was not found.")
        rows = [row for row in rows if row[column] == options.get("filter_value", "")]
        scope += "; filter " + column + " = " + options.get("filter_value", "")
    if options.get('scope_note'):
        scope += '; ' + options['scope_note']
    if not rows:
        raise ValueError("No records extracted. Empty results cannot replace an accepted dataset.")
    return {"headers": headers, "rows": rows, "scope": scope}

if __name__ == "__main__":
    try:
        result = extract(Path(sys.argv[1]), sys.argv[2], json.load(sys.stdin))
        encoded = json.dumps(result, ensure_ascii=False).encode('utf-8')
        if len(sys.argv) > 3:
            Path(sys.argv[3]).write_bytes(encoded)
        else:
            sys.stdout.buffer.write(encoded)
    except ValueError as error:
        sys.stdout.buffer.write(json.dumps({"error": str(error)}).encode("utf-8"))
        sys.exit(1)
    except Exception:
        sys.stdout.buffer.write(json.dumps({"error": "Source could not be parsed. Check its format, worksheet/path and extraction dependencies."}).encode("utf-8"))
        sys.exit(1)
