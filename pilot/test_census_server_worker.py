import hashlib
from contextlib import closing
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import openpyxl
from census_server_worker import run, workbook_rows


class CensusWorkerTests(unittest.TestCase):
    def test_ocr_resource_pause_resumes_without_error_backoff(self):
        from ocr_civic_pdf_pages import ResourceWait
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); (root/'queue').mkdir()
            (root/'queue/test.json').write_text(json.dumps({'kind':'pdf_ocr', 'package':'test.pdf',
                'sha256':'a'*64,'pages':[1],'source_url':'https://censusindia.gov.in/nada/test'}))
            with patch('census_server_worker.resources_ok',return_value=True), patch(
                    'ocr_civic_pdf_pages.extract',side_effect=ResourceWait('Waiting for disk/RAM reserve')):
                run(root)
            with closing(sqlite3.connect(root/'census-review.sqlite')) as db:
                self.assertEqual(db.execute('SELECT status,error,retry_after FROM jobs').fetchone(),('pending',None,0))
            self.assertEqual(json.loads((root/'status.json').read_text())['state'],'waiting_for_resources')
            with patch('census_server_worker.resources_ok',return_value=True), patch('ocr_civic_pdf_pages.extract') as extract:
                run(root)
                extract.assert_called_once()
            with closing(sqlite3.connect(root/'census-review.sqlite')) as db:
                self.assertEqual(db.execute('SELECT status FROM jobs').fetchone()[0],'complete')

    def test_invalid_styles_fallback_keeps_raw_cells_and_formula(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'raw.xlsx'
            book = openpyxl.Workbook()
            sheet = book.active
            sheet['A1'] = 'Code'
            sheet['C1'] = 'Value'
            sheet['A2'] = '0012'
            sheet['C2'] = 3
            sheet['D2'] = '=C2+1'
            book.save(path)
            with patch('openpyxl.load_workbook', side_effect=TypeError("expected <class 'openpyxl.styles.fills.Fill'>")):
                rows = list(workbook_rows(path))
            self.assertEqual([row[1] for row in rows], [1, 2])
            self.assertEqual(rows[1][2], ['0012', None, '3', '=C2+1'])
            self.assertEqual(rows[1][3][2]['type'], 'n')
            self.assertEqual(rows[1][3][3]['formula'], 'C2+1')
            self.assertEqual(rows[1][3][2]['reader'], 'raw_ooxml_invalid_styles')

    def test_lgd_sparse_cells_keep_source_indices_and_formula(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'lgd.zip'
            body = '''<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="Villages"><Table><Row ss:Index="4"><Cell><Data ss:Type="String">0012</Data></Cell><Cell ss:Index="3" ss:Formula="=R4C1"><Data ss:Type="Number">12</Data></Cell></Row></Table></Worksheet></Workbook>'''
            with zipfile.ZipFile(path, 'w') as archive:
                archive.writestr('villages.xls', body)
            rows = list(workbook_rows(path))
            self.assertEqual(rows[0][:3], ('villages.xls/Villages',4,['0012',None,'12']))
            self.assertEqual(rows[0][3][2]['formula'], '=R4C1')

    def fixture(self, root, corrupt=False):
        (root / 'queue').mkdir()
        (root / 'packages').mkdir()
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(['Area', 'Count'])
        sheet.append(['Example', 3])
        sheet.append(['Formula', '=B2'])
        stream = io.BytesIO()
        book.save(stream)
        body = stream.getvalue()
        raw_hash = hashlib.sha256(body).hexdigest()
        extracted = b'{"rows":[]}'
        sources = [dict(key=key, sha256=raw_hash, extracted_sha256=hashlib.sha256(extracted).hexdigest(),
                        source_url='https://censusindia.gov.in/example', options={'scope_note':'unverified'})
                   for key in ('example-total', 'example-rural')]
        package = root / 'packages' / 'test.zip'
        with zipfile.ZipFile(package, 'w') as archive:
            archive.writestr('manifest.json', json.dumps({'version':1, 'sources':sources}))
            for source in sources:
                archive.writestr(source['key'] + '.xlsx', body if not corrupt else body + b'bad')
                archive.writestr(source['key'] + '.json', extracted)
        (root / 'queue' / 'test.json').write_text(json.dumps({'package':'test.zip','sha256':hashlib.sha256(package.read_bytes()).hexdigest()}))
        return raw_hash

    @patch('census_server_worker.resources_ok', return_value=True)
    def test_resume_deduplicates_original_and_preserves_formula(self, _):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source_hash = self.fixture(root)
            run(root)
            run(root)
            db = sqlite3.connect(root / 'census-review.sqlite')
            self.assertEqual(db.execute('SELECT count(*) FROM source_references').fetchone()[0], 2)
            self.assertEqual(db.execute('SELECT count(*) FROM raw_rows').fetchone()[0], 3)
            self.assertIn('=B2', db.execute('SELECT cells_json FROM raw_rows WHERE source_row=3').fetchone()[0])
            # Simulate interrupted source, then confirm restart replaces partial rows.
            db.execute("UPDATE workbooks SET status='extracting' WHERE sha256=?", (source_hash,))
            db.execute("UPDATE jobs SET status='running'")
            db.commit()
            db.close()
            run(root)
            with closing(sqlite3.connect(root / 'census-review.sqlite')) as db:
                self.assertEqual(db.execute('SELECT count(*) FROM raw_rows').fetchone()[0], 3)
                self.assertEqual(db.execute('PRAGMA integrity_check').fetchone()[0], 'ok')

    @patch('census_server_worker.resources_ok', return_value=True)
    def test_rejects_corrupt_original_without_extracted_rows(self, _):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            self.fixture(root, corrupt=True)
            run(root)
            with closing(sqlite3.connect(root / 'census-review.sqlite')) as db:
                self.assertEqual(db.execute('SELECT count(*) FROM raw_rows').fetchone()[0], 0)
                self.assertEqual(db.execute('SELECT status FROM jobs').fetchone()[0], 'error')

    @patch('census_server_worker.resources_ok', return_value=True)
    def test_validates_original_cells_and_catches_wrong_values(self, _):
        for wrong_value in (False, True):
            with self.subTest(wrong_value=wrong_value), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                source_hash = self.fixture(root)
                with zipfile.ZipFile(root / 'packages/test.zip') as archive:
                    body = archive.read('example-total.xlsx')
                package = root / 'packages/originals.zip'
                with zipfile.ZipFile(package, 'w') as archive:
                    archive.writestr('123/' + source_hash + '.xlsx', body)
                review_path = root / 'review.sqlite'
                with closing(sqlite3.connect(review_path)) as review:
                    review.executescript('''CREATE TABLE sources(id,source_sha256);
                        CREATE TABLE sheets(source_id,number,name,headers_json,expected_rows);
                        CREATE TABLE source_rows(source_id,sheet_number,source_row,cells_json);''')
                    review.execute('INSERT INTO sources VALUES (?,?)', ('test', source_hash))
                    review.execute('INSERT INTO sheets VALUES (?,?,?,?,?)', ('test',0,'Sheet','["Area","Count"]',2))
                    review.executemany('INSERT INTO source_rows VALUES (?,?,?,?)', [
                        ('test',0,2,json.dumps(['Example',9 if wrong_value else 3])),
                        ('test',0,3,json.dumps(['Formula','=B2']))])
                    review.commit()
                (root / 'queue/test.json').write_text(json.dumps({'kind':'validate_1991',
                    'package':package.name, 'sha256':hashlib.sha256(package.read_bytes()).hexdigest(),
                    'review_db':str(review_path), 'review_sha256':hashlib.sha256(review_path.read_bytes()).hexdigest()}))
                run(root)
                with closing(sqlite3.connect(root / 'census-review.sqlite')) as db:
                    self.assertEqual(db.execute('SELECT status FROM jobs').fetchone()[0], 'error' if wrong_value else 'complete')
                    self.assertEqual(db.execute('SELECT count(*) FROM source_validations').fetchone()[0], 0 if wrong_value else 1)


if __name__ == '__main__':
    unittest.main()
