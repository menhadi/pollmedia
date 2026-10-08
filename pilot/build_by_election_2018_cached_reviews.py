"""Preserve formula cells while adding checksum-bound workbook cache reviews."""
import copy
import hashlib
import json
from pathlib import Path
import re
import tempfile
import zipfile

import openpyxl
from preserve_archive_json import package

ROOT = Path(__file__).resolve().parents[1]
NAME = 'pollmedia-by-election-2018-cached-reviews-20261009'
PREFIX = 'election-by-elections/structured/'
EDITION = '29b120100d2e2ac505c50122'
SOURCE_SHA = '71ca0ec9bc18c73f420800649703eff4f5d7c10563da49d9f72e95a53aeb8e3e'
INDEX_SHA = '2c2a861be15b49bbd77f018063f9523131371154ef8254201a2b965d457dbfe8'
URL = 'https://old.eci.gov.in/ByeElection/2018/ByeElection(Jan-June2018).xlsx'
TARGETS = {
    'd57a8c5fb4c4525889fa26b5': ('Gorakhpur PC UP', '610123892ebcc29e767a0b3099a6b1e48fc946dcc12afbd8b6cb109aaaba65da', 1955461, 934441, 925989),
    '0ddcb16492a341891faedeec': ('1-Nagaland PC', '6fafb93b4832a9ed11b31e8cc38cd5ac20e8d902c46e8fad3e72dc1242f8ec24', 1197436, 1018842, 1014664),
}


def digest(body):
    return hashlib.sha256(body).hexdigest()


def revised_files(root=ROOT):
    folder = root / 'application/storage/app/private'
    index_bytes = (folder / (PREFIX + 'index.json')).read_bytes()
    if digest(index_bytes) != INDEX_SHA:
        raise ValueError('Prior by-election index differs')
    index = json.loads(index_bytes)
    revised = copy.deepcopy(index)
    source = folder / 'election-by-elections' / EDITION / (SOURCE_SHA + '.xlsx')
    if source.is_symlink() or digest(source.read_bytes()) != SOURCE_SHA:
        raise ValueError('Official workbook checksum differs')
    formulas = openpyxl.load_workbook(source, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(source, read_only=True, data_only=True)
    originals, additions, reviews = {}, {}, []
    try:
        for rid, (sheet, prior, electors, polled, valid) in TARGETS.items():
            matches = [r for r in revised['records'] if r['id'] == rid]
            if len(matches) != 1 or matches[0]['sha256'] != prior:
                raise ValueError('Target index identity differs')
            entry = matches[0]
            old_path = PREFIX + entry['file']
            body = (folder / old_path).read_bytes()
            if digest(body) != prior:
                raise ValueError('Prior structured record differs')
            record = json.loads(body)
            if (record['id'] != rid or record['year'] != 2018 or record['kind'] != 'pc'
                    or record['edition'] != EDITION or record['source_sha256'] != SOURCE_SHA
                    or record['source_url'] != URL or 'source_review' in record):
                raise ValueError('Structured source identity differs')
            f, c = formulas[sheet], cached[sheet]
            candidates = []
            for candidate in record['candidates']:
                row = candidate['source_row']
                raw, value = f.cell(row, 4).value, c.cell(row, 4).value
                if raw != candidate['raw_votes'] or type(value) is not int or value < 0:
                    raise ValueError('Candidate vote cell differs')
                if isinstance(raw, str):
                    if (not re.fullmatch(r'=[+-]?\d+(?:[+-]\d+)*', raw)
                            or sum(int(v) for v in re.findall(r'[+-]?\d+', raw[1:])) != value):
                        raise ValueError('Cached formula value is not independently reconciled')
                elif raw != value:
                    raise ValueError('Literal vote differs')
                candidates.append({'name': candidate['name'], 'votes': value, 'cell': f'D{row}', 'original': raw})
            if (c['F18'].value != electors or c['F30'].value != polled or c['F31'].value != valid
                    or sum(x['votes'] for x in candidates) != valid
                    or c.cell(41 + len(candidates), 4).value != valid
                    or c['F12'].value != len(candidates)
                    or valid + c['F32'].value + c['F33'].value != polled
                    or c['F25'].value + c['F27'].value != polled or c['F28'].value != polled):
                raise ValueError('Source totals do not reconcile')
            notes = ['Saved workbook vote values agree with the literal formulas, candidate count and valid-vote total. Original formula cells, missing extracted votes and warnings are preserved.']
            if rid.startswith('d57'):
                if str(c['A23'].value) != '2018-11-03 00:00:00' or c['B23'].value != '14-3-2018' or c['C23'].value != '14-3-2018':
                    raise ValueError('Gorakhpur date discrepancy differs')
                notes.append('The polling cell is an Excel date of 3 November 2018, while counting and declaration cells say 14 March 2018. The source dates conflict; no corrected polling date is inferred.')
            else:
                if [c[cell].value for cell in ['A23', 'B23', 'C23']] != ['28-5-2018', '31-5-2018', '31-5-2018']:
                    raise ValueError('Nagaland dates differ')
                notes.append('Source dates: polling 28 May 2018; counting and declaration 31 May 2018. Party names remain as printed in the original rows.')
            after = copy.deepcopy(record)
            after['source_review'] = {'sheet': sheet, 'source_sha256': SOURCE_SHA, 'prior_sha256': prior,
                'candidates': candidates, 'totals': {'electors': electors, 'votes_polled': polled, 'valid_candidate_votes': valid}, 'notes': notes}
            new_body = json.dumps(after, ensure_ascii=False, indent=2).encode()
            new_sha = digest(new_body)
            entry['file'] = rid + '-' + new_sha[:16] + '.json'
            entry['sha256'] = new_sha
            originals[old_path] = body
            additions[PREFIX + entry['file']] = new_body
            reviews.append({'id': rid, 'prior_path': old_path, 'prior_sha256': prior, 'path': PREFIX + entry['file'], 'sha256': new_sha})
    finally:
        formulas.close()
        cached.close()
    for old, new in zip(index['records'], revised['records'], strict=True):
        if old['id'] not in TARGETS and old != new:
            raise ValueError('Unrelated index entry changed')
    return index_bytes, json.dumps(revised, ensure_ascii=False, indent=2).encode(), originals, additions, reviews


def import_script():
    return '''#!/usr/bin/env bash
set -euo pipefail
cd "$(dirname "$0")"
exec 9> /home/pollmedia/tmp/.pollmedia-election-release.lock
flock -n 9 || { echo 'Election import already running' >&2; exit 1; }
sha256sum -c SHA256SUMS
artisan=/home/pollmedia/app/application/artisan
test -f "$artisan"
test -f /home/pollmedia/app/application/resources/views/by-election-source-review.blade.php
if pgrep -af 'archive:import-json|archive:index-constituencies|[e]lection.*[o]cr'; then
    echo 'Another election import, index or OCR process is active' >&2; exit 1
fi
check_disk() {
    available=$(df -Pk /home/pollmedia/app | awk 'NR==2 {print $4}')
    test "$available" -ge 10551296 || { echo 'Cannot maintain 10 GiB reserve plus package headroom' >&2; exit 1; }
}
check_disk
php8.4 PREFLIGHT.php
for file in snapshot-*.zip addition-*.zip; do
    check_disk
    checksum=$(sha256sum "$file" | awk '{print $1}')
    php8.4 "$artisan" archive:import-json "$PWD/$file" "--sha256=$checksum" --check
    php8.4 "$artisan" archive:import-json "$PWD/$file" "--sha256=$checksum"
done
check_disk
php8.4 PREFLIGHT.php
file=correction-index.zip
checksum=$(sha256sum "$file" | awk '{print $1}')
php8.4 "$artisan" archive:import-json "$PWD/$file" "--sha256=$checksum" --check --allow-revision
php8.4 "$artisan" archive:import-json "$PWD/$file" "--sha256=$checksum" --allow-revision
check_disk
echo 'Two by-election source reviews imported; originals retained. Not accepted contests.'
'''


def preflight():
    return '''<?php
require '/home/pollmedia/app/application/vendor/autoload.php';
$app = require '/home/pollmedia/app/application/bootstrap/app.php';
$app->make(\\Illuminate\\Contracts\\Console\\Kernel::class)->bootstrap();
$audit = json_decode(file_get_contents(__DIR__.'/AUDIT.json'), true, 512, JSON_THROW_ON_ERROR);
$expected = $audit['required_live'];
foreach ($expected as $path => $hashes) {
    if (is_file(storage_path('app/private/'.$path))) { throw new RuntimeException('Local JSON would shadow database: '.$path); }
    $row = \\Illuminate\\Support\\Facades\\DB::table('archive_json_files')->where('path_hash', hash('sha256', $path))->first();
    if (!$row || $row->path !== $path || !in_array($row->sha256, $hashes, true) || !hash_equals($row->sha256, hash('sha256', $row->body))) {
        throw new RuntimeException('Required live checksum differs: '.$path);
    }
}
foreach ($audit['additions'] as $path => $sha) {
    if (is_file(storage_path('app/private/'.$path))) { throw new RuntimeException('Local review JSON shadows database: '.$path); }
    $row = \\Illuminate\\Support\\Facades\\DB::table('archive_json_files')->where('path_hash', hash('sha256', $path))->first();
    if ($row && ($row->path !== $path || $row->sha256 !== $sha || hash('sha256', $row->body) !== $sha)) { throw new RuntimeException('Conflicting review JSON: '.$path); }
}
echo "PASS: exact live predecessors or completed revision verified.\\n";
'''


def build(root=ROOT):
    old, new, originals, additions, reviews = revised_files(root)
    output = root / 'exports' / (NAME + '.zip')
    if output.exists() or output.with_suffix('.sha256').exists():
        raise FileExistsError(output)
    snapshot = PREFIX + 'index-' + INDEX_SHA + '.json'
    audit = {'reviews': reviews, 'prior_index_sha256': INDEX_SHA, 'new_index_sha256': digest(new),
             'required_live': {PREFIX + 'index.json': [INDEX_SHA, digest(new)], **{p: [digest(b)] for p, b in originals.items()}},
             'additions': {p: digest(b) for p, b in additions.items()}, 'scope': 'Two reviewed by-election workbooks; not accepted contests.'}
    with tempfile.TemporaryDirectory(dir=output.parent, prefix='by-election-review-') as tmp:
        staged = Path(tmp) / 'staged'
        files = {snapshot: old, PREFIX + 'index.json': new, **originals, **additions}
        for path, body in files.items():
            target = staged / path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(body)
        packages = []
        ordered = [('snapshot-index', snapshot), *[(f'snapshot-record-{i}', p) for i, p in enumerate(originals)],
                   *[(f'addition-record-{i}', p) for i, p in enumerate(additions)], ('correction-index', PREFIX + 'index.json')]
        for name, path in ordered:
            z = Path(tmp) / (name + '.zip')
            revision = name == 'correction-index'
            package(staged, z, 'election-by-elections', 0, 1, [path], INDEX_SHA if revision else None, snapshot if revision else None)
            packages.append(z)
        content = {'AUDIT.json': json.dumps(audit, indent=2).encode(), 'PREFLIGHT.php': preflight().encode()}
        content.update({p.name: p.read_bytes() for p in packages})
        with zipfile.ZipFile(output, 'x', compression=zipfile.ZIP_STORED) as bundle:
            for name, body in content.items():
                bundle.writestr(name, body)
            bundle.writestr('SHA256SUMS', ''.join(f'{digest(b)}  {n}\n' for n, b in content.items()))
            bundle.writestr('IMPORT.sh', import_script())
    sha = digest(output.read_bytes())
    output.with_suffix('.sha256').write_bytes(f'{sha}  {output.name}\n'.encode())
    return {'bundle': str(output), 'sha256': sha, 'prior_index_sha256': INDEX_SHA, 'new_index_sha256': digest(new)}


if __name__ == '__main__':
    print(json.dumps(build()))
