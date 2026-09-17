"""Read the official UP 2017 workbook pair, withholding missing or inconsistent summaries."""
import json
import re
import sys
from pathlib import Path
import openpyxl
from extract_election_import import candidate, integer, name, safe_zip


def normal(value):
    return re.sub(r"[^a-z0-9]", "", name(value).lower())


def extract_state(detail, summary):
    for path in [detail, summary]:
        with safe_zip(path):
            pass
    groups = {}
    book = openpyxl.load_workbook(detail, read_only=True, data_only=False)
    try:
        sheet = book['DetailedResult']
        for index, row in enumerate(sheet.values, 1):
            if index == 2 and row[:3] != ('Constituency No.', 'Constituency Name', 'Candidate Name'):
                raise ValueError('2017 detailed header changed.')
            if index <= 2 or all(v is None for v in row):
                continue
            if index > 20000 or type(row[0]) is not int or not 1 <= row[0] <= 403:
                raise ValueError('Unexpected constituency code or report dimensions.')
            groups.setdefault(row[0], []).append((index, row))
    finally:
        book.close()
    if sorted(groups) != list(range(1, 404)):
        raise ValueError('Detailed report must contain exactly 403 seats.')
    book = openpyxl.load_workbook(summary, read_only=True, data_only=False)
    try:
        sheets = {}
        for title in book.sheetnames:
            match = re.fullmatch(r'(\d+)\s+(.+)', title.strip())
            if not match or int(match[1]) in sheets:
                raise ValueError('Unrecognized or duplicate summary sheet identity.')
            sheets[int(match[1])] = title
        results = []
        for code, entries in sorted(groups.items()):
            seat = name(entries[0][1][1])
            record = dict(code=code, name=seat)
            try:
                if code not in sheets:
                    raise ValueError('Official 2017 summary workbook and PDF omit this constituency; total votes polled cannot be verified.')
                sheet = book[sheets[code]]
                cells = {cell.coordinate: cell.value for row in sheet for cell in row if cell.value is not None}
                if normal(cells.get('B2')) != normal(seat):
                    raise ValueError('Detailed and summary constituency names differ.')
                for cell, label in {'B12':'General(other than OVERSEAS)', 'B13':'OVERSEAS', 'B14':'Service', 'B18':'General(other than OVERSEAS)', 'B19':'OVERSEAS', 'B21':'Postal', 'B26':'Rejected Votes (Postal)', 'B27':'Votes Not Retrieved From EVM'}.items():
                    if normal(cells.get(cell)) != normal(label):
                        raise ValueError('Summary labels changed at ' + cell)
                electors = sum(integer(cells[c+str(r)]) for r in [12,13,14] for c in ['D','F','H'])
                polled = sum(integer(cells[c+str(r)]) for r in [18,19] for c in ['D','F','H']) + integer(cells['J21'])
                rejected, unretrieved = integer(cells['C26']), integer(cells['C27'])
                candidates = []
                for rank, (source, row) in enumerate(entries, 1):
                    if normal(row[1]) != normal(seat):
                        raise ValueError('Constituency names differ within the detailed report.')
                    if integer(row[10]) != electors:
                        raise ValueError(f'Elector totals conflict: detailed report {integer(row[10])}; summary components {electors}.')
                    candidates.append(candidate(rank, row[2], row[6], row[7], row[8], row[9]))
                recorded = sum(r['votes'] for r in candidates)
                if any(integer(row[11]) != recorded for _,row in entries) or recorded != polled - rejected - unretrieved:
                    raise ValueError('Recorded votes do not reconcile with the detailed and summary reports.')
                record['payload'] = dict(year=2017,code=code,name=seat,candidates=candidates,electors=electors,votes_polled=polled,valid_candidate_votes=sum(r['votes'] for r in candidates if not r['is_nota']),rejected_postal_votes=rejected,unretrieved_evm_votes=unretrieved,source_locator=f'DetailedResult rows {entries[0][0]}-{entries[-1][0]}; summary sheet {sheets[code]}, literal elector/voter components D/F/H12-14,18-19; J21; C26-C27. Formula cells are not evaluated.')
            except (ValueError,KeyError,TypeError) as error:
                record['error'] = str(error) if isinstance(error,ValueError) else 'Missing or invalid summary component.'
            results.append(record)
        return results
    finally:
        book.close()


if __name__ == '__main__':
    try:
        sys.stdout.buffer.write(json.dumps(extract_state(Path(sys.argv[1]),Path(sys.argv[2])),ensure_ascii=False).encode('utf-8'))
    except Exception as error:
        sys.stdout.buffer.write(json.dumps({'error':str(error)}).encode('utf-8'))
        sys.exit(1)
