"""Review dash placeholders without replacing preserved historical source records."""
import copy
from io import BytesIO
import json
import zipfile

import openpyxl

from build_by_election_2018_cached_reviews import ROOT, PREFIX, build_package, digest

NAME = 'pollmedia-by-election-historical-placeholder-reviews-20261009'
PREDECESSOR = 'pollmedia-by-election-1996-1999-source-notes-20261009.zip'
PACKAGE_SHA = '37ac0a4178e9b78e71ccf564a5d124125387915b2cbfb0f3778e4a5b0928dab9'
INDEX_SHA = 'd744a0b9dd157fa6b5a6e73080c37579c584bf97cc51b82ac860d55e7b0c6a1c'
SOURCE_SHA = 'd2471c1cb7fe962eeb56791f2025463656696fc7d0a08ffe4a254be089bfa680'
TARGETS = {
    '2291b041ec8724b3a8c51cde': ('42e55d328156254bf29001c8392eff13177e21cb219cefbce756efaed80b452a', 2),
    'fb074ed2d5aac3e343273f22': ('cbc54a1ff0c0d24adb3e9a1fc8209d1ce9995825563dcbe0b801f903db9ed5ba', 4),
    'f3fca4aadc25485d912a5bd1': ('7fa1df273b5c2d381d1b781e48b73898161b554262e9962008e12b909253c62b', 1),
}


def reviewed_record(record, source_body, prior, removed_count):
    if digest(source_body) != SOURCE_SHA or record['source_sha256'] != SOURCE_SHA:
        raise ValueError('Official workbook checksum differs')
    book = openpyxl.load_workbook(BytesIO(source_body), read_only=True, data_only=False)
    try:
        for candidate in record['candidates']:
            row = candidate['source_row']
            cells = list(next(book[candidate['table']].iter_rows(min_row=row, max_row=row, values_only=True)))
            if cells != candidate['source_cells']:
                raise ValueError('Original candidate source cells differ')
            name_col = 6 if candidate['role'] == 'reported_winner' else 9
            if candidate['name'] != ' '.join(str(cells[name_col]).split()):
                raise ValueError('Original candidate name differs from source')
    finally:
        book.close()
    placeholders = [c for c in record['candidates'] if c['name'] == '-']
    if len(placeholders) != removed_count or any(c['votes'] is not None for c in placeholders):
        raise ValueError('Expected nonnumeric dash placeholders differ')
    revised = copy.deepcopy(record)
    revised['original_candidate_rows'] = copy.deepcopy(record['candidates'])
    revised['original_candidate_count'] = record['candidate_count']
    revised['candidates'] = [c for c in revised['candidates'] if c['name'] != '-']
    revised['candidate_count'] = len(revised['candidates'])
    revised['notes'].append(
        f'Source review: {removed_count} dash-only name cells were previously counted as candidates. '
        'They are now excluded from the displayed candidate list; all named source rows, original '
        'candidate rows and earlier warnings are preserved. Earlier additional-member warnings may '
        'include these placeholders. This historical summary does not establish complete candidate '
        'coverage, a single-seat winner, ordinary turnout or margin, or an accepted contest.')
    revised['historical_placeholder_review'] = {
        'source_sha256': SOURCE_SHA, 'prior_sha256': prior,
        'excluded_placeholders': copy.deepcopy(placeholders),
    }
    return revised


def revised_files(root=ROOT):
    path = root/'exports'/PREDECESSOR
    if digest(path.read_bytes()) != PACKAGE_SHA:
        raise ValueError('Predecessor package differs')
    with zipfile.ZipFile(path) as outer, zipfile.ZipFile(BytesIO(outer.read('correction-index.zip'))) as inner:
        old = inner.read(PREFIX+'index.json')
    if digest(old) != INDEX_SHA:
        raise ValueError('Predecessor index differs')
    index = json.loads(old)
    originals, additions, reviews = {}, {}, []
    folder = root/'application/storage/app/private'
    for rid, (prior, removed_count) in TARGETS.items():
        entries = [e for e in index['records'] if e['id'] == rid]
        if len(entries) != 1 or entries[0]['sha256'] != prior:
            raise ValueError('Prior index differs')
        entry = entries[0]
        prior_path = PREFIX+entry['file']
        body = (folder/prior_path).read_bytes()
        if digest(body) != prior:
            raise ValueError('Prior extraction differs')
        record = json.loads(body)
        source = folder/'election-by-elections'/record['edition']/record['source_file']
        if source.is_symlink():
            raise ValueError('Source must be a regular archive file')
        revised = reviewed_record(record, source.read_bytes(), prior, removed_count)
        new_body = json.dumps(revised, ensure_ascii=False, indent=2).encode()
        sha = digest(new_body)
        entry.update(file=rid+'-'+sha[:16]+'.json', sha256=sha,
                     candidate_count=revised['candidate_count'], note_count=len(revised['notes']))
        new_path = PREFIX+entry['file']
        originals[prior_path], additions[new_path] = body, new_body
        reviews.append({'id': rid, 'prior_path': prior_path, 'prior_sha256': prior, 'path': new_path, 'sha256': sha})
    return old, json.dumps(index, ensure_ascii=False, indent=2).encode(), originals, additions, reviews


def build(root=ROOT):
    return build_package(root, NAME, *revised_files(root))


if __name__ == '__main__':
    print(json.dumps(build()))
